"""db/toeic_rc.db를 열람용 엑셀(db/toeic_rc_items.xlsx)로 내보낸다. Part별 시트, 유형 순 정렬."""
import json
import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
con = sqlite3.connect(os.path.join(ROOT, "db", "toeic_rc.db"))
wb = Workbook()
wb.remove(wb.active)
cols = ["item_id", "세트", "대분류", "유형 태그", "난이도", "트랙", "문제", "A", "B", "C", "D", "정답", "핵심 규칙", "해설", "상태"]
sql = """
SELECT i.item_id, COALESCE(i.set_id, ''), t.category_ko, i.type_tag, i.difficulty_expert, i.target_track, i.stem,
  (SELECT body FROM item_options WHERE item_id=i.item_id AND label='A'),
  (SELECT body FROM item_options WHERE item_id=i.item_id AND label='B'),
  (SELECT body FROM item_options WHERE item_id=i.item_id AND label='C'),
  (SELECT body FROM item_options WHERE item_id=i.item_id AND label='D'),
  i.answer, i.key_point, i.explanation_ko, i.status
FROM items i JOIN tag_codes t ON t.code = i.type_tag WHERE i.part = ?
ORDER BY i.category, i.type_tag, i.item_id
"""
for part in (5, 6, 7):
    ws = wb.create_sheet(f"Part {part}")
    ws.append(cols)
    for row in con.execute(sql, (part,)):
        ws.append(list(row))
    for c in ws[1]:
        c.font = Font(bold=True)
    for letter, w in zip("ABCDEFGHIJKLMNO", [14, 14, 10, 18, 7, 16, 60, 22, 22, 22, 22, 6, 40, 60, 14]):
        ws.column_dimensions[letter].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
ws = wb.create_sheet("지문 (Part 6·7)")
ws.append(["set_id", "순서", "지문 유형", "제목", "단어 수", "본문"])
for row in con.execute("SELECT set_id, seq, passage_type, title, word_count, body FROM passages ORDER BY set_id, seq"):
    ws.append(list(row))
ws.column_dimensions["F"].width = 100
for r in ws.iter_rows(min_row=2):
    r[5].alignment = Alignment(wrap_text=True, vertical="top")
ws = wb.create_sheet("유형별 집계")
ws.append(["Part", "대분류", "유형 태그", "문항 수", "평균 난이도"])
for p, cat, cat_ko, tag, n, avg in con.execute("SELECT * FROM v_tag_summary"):
    ws.append([p, cat_ko, tag, n, avg])
ws = wb.create_sheet("어휘")
ws.append(["단어", "품사", "뜻", "수준", "나온 문항"])
for row in con.execute("""SELECT v.word, v.pos, v.meaning_ko, v.level, GROUP_CONCAT(iv.item_id, ', ')
    FROM vocab v LEFT JOIN item_vocab iv ON iv.word=v.word AND iv.pos=v.pos GROUP BY v.word, v.pos ORDER BY v.word"""):
    ws.append(list(row))
ws = wb.create_sheet("어휘 카드")
ws.append(["word_id", "레벨", "표제어", "품사", "뜻", "예문", "번역", "콜로케이션", "파생어", "유형 태그", "상태"])
for wid, lv, lemma, pos, mean, ex, col, der, tags, st in con.execute(
        "SELECT word_id, level, lemma, pos, meaning_ko, examples, collocations, derivatives, type_tags, status FROM word_cards ORDER BY word_id"):
    ex = json.loads(ex)
    ws.append([wid, lv, lemma, pos, mean, ex[0]["en"], ex[0]["ko"], " / ".join(json.loads(col)),
               ", ".join(d["form"] + "(" + d["pos"] + ")" for d in json.loads(der)), ", ".join(json.loads(tags)), st])
for letter, w in zip("ABCDEFGHIJK", [12, 6, 16, 6, 24, 50, 40, 40, 36, 20, 12]):
    ws.column_dimensions[letter].width = w
ws = wb.create_sheet("배치고사 폼")
ws.append(["form_id", "단계", "순서", "item_id", "Part", "유형 태그", "난이도"])
for row in con.execute("""SELECT fs.form_id, fs.stage, fs.seq, fs.item_id, i.part, i.type_tag, i.difficulty_expert
    FROM form_slots fs JOIN items i ON i.item_id = fs.item_id ORDER BY fs.form_id, fs.stage, fs.seq"""):
    ws.append(list(row))
wb.save(os.path.join(ROOT, "db", "toeic_rc_items.xlsx"))
print("ok")
